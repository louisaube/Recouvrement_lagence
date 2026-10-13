#!/usr/bin/env python3
"""analyse_fine.py : rétro-analyse d'un classeur xlsx (enveloppe, plan, graphe, sens, comportement, croisements).
  python analyse_fine.py modele.xlsx --sorties "'Compte de résultat'!D57" --md rapport.md --json carte.json
  --effets : recalcul cellule par cellule dans LibreOffice (effet chiffré, seuils, absorbeurs, sélecteurs par valeur)
  --cartes DOSSIER : carte d'ensemble et carte des couloirs (PNG) ; --libelles lib.json ; --verifier-cache
Sorties par défaut : rôles output de l'onglet Map, sinon puits du graphe (cône approximatif)."""
import argparse, json, logging, math, os, re, shutil, subprocess, tempfile, time, zipfile
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from datetime import date, datetime
from functools import lru_cache
from pathlib import Path

import networkx as nx
from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.utils import column_index_from_string as CI, get_column_letter as GL, range_boundaries
from openpyxl.utils.datetime import to_excel
from openpyxl.worksheet.formula import ArrayFormula

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
OPAQUES = {"INDIRECT", "OFFSET"}
SI = {"SUMIF", "AVERAGEIF", "COUNTIF", "SUMIFS", "AVERAGEIFS", "COUNTIFS", "MAXIFS", "MINIFS"}
RECH = {"VLOOKUP", "HLOOKUP", "MATCH", "XLOOKUP", "LOOKUP"}
CHOIX = {"IF", "IFS", "SWITCH"}
FORME = {"COUNT", "COUNTA", "COUNTBLANK", "ROWS", "COLUMNS", "ROW", "COLUMN", "ISNUMBER", "ISTEXT", "ISBLANK",
         "ISERROR", "ISERR", "ISNA", "ISLOGICAL", "ISNONTEXT", "ISFORMULA", "ISREF"}   # lisent la présence, pas la valeur
TESTS = {"ISNUMBER": lambda v: nombre(v) is not None, "ISTEXT": lambda v: isinstance(v, str) and v not in ERR,
         "ISBLANK": lambda v: v is None, "ISERROR": lambda v: v in ERR, "ISNA": lambda v: v == "#N/A",
         "ISLOGICAL": lambda v: isinstance(v, bool), "ISNONTEXT": lambda v: not isinstance(v, str) or v in ERR}
CMP = {"=", "<>", "<", ">", "<=", ">="}
ERR = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}
TRIVIAUX = {0, 0.5, 1, 2, 3, 4, 5, 6, 7, 10, 12, 24, 30, 31, 52, 60, 100, 365, 366, 1000, 1e6}
MOTS_VIDES = {"de", "du", "des", "la", "le", "les", "et", "vs", "à", "d", "l", "en", "par"}
INCONNU = object()
RE_A1 = re.compile(r"(?<![A-Za-z0-9_.])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![A-Za-z0-9_(])")
RE_AN = re.compile(r"(?<!\d)(19[5-9]\d|20\d\d|21[0-4]\d)(?!\d)")
RE_LAG = re.compile(r"ouverture|cl[ôo]ture|solde|report|opening|closing|pr[ée]c[ée]dent|cumul|stock|variation|n-1|m-1", re.I)
RE_CHK = re.compile(r"\bchecks?\b|contr[ôo]les?\b|v[ée]rif|doit [êe]tre|coh[ée]rence", re.I)
RE_DOC = re.compile(r"lisez|read ?me|notice|guide|documentation|sommaire|mode d.emploi|[àa] propos", re.I)
RE_SYNTH = re.compile(r"total|cumul|moy|ytd|ann[ée]e|annuel|[àa] fin|variation|[ée]cart|ratio|%|taux|marge|ebitda|"
                      r"r[ée]sultat|solde", re.I)
RE_CALC = re.compile(r"(?<![<>=])=(?!=)|×|÷|\*|\s[xX]\s|\bpour\s+\d|\bpar\s+\d")
RE_LIGNE = re.compile(r"\b(?:lignes?|l\.)\s*(\d{1,9}(?:\s*(?:,|et|à|-)\s*\d{1,9})*)", re.I)
ESP = " " + chr(0xA0) + chr(0x202F)                     # espaces des milliers : normale, insécable, fine
RE_NB = re.compile(r"\d+(?:[" + ESP + r"]\d{3})*(?:[.,]\d+)?")
RE_SOMME = re.compile(r"^=\s*(SUM\(|SUBTOTAL\(\s*(?:9|109)\s*,)[^()]*\)\s*$", re.I)
SEV = ["critique", "majeure", "à vérifier"]


def formule(v):
    t = v if isinstance(v, str) else getattr(v, "text", None)            # ArrayFormula
    return t if isinstance(t, str) and t.startswith("=") else None


def nombre(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return float(to_excel(v)) if isinstance(v, (datetime, date)) else None


def lire(ws, l, c):
    x = ws._cells.get((l, c))
    return None if x is None else x.value


def plat(m):
    return [x for r in m for x in r]


def slug(s, n=28):
    s = str(s or "x").lower().translate(str.maketrans("éèêëàâîïôöûùç", "eeeeaaiioouuc"))
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")[:n] or "x"


@lru_cache(maxsize=None)
def pos(n):
    """('Feuille', ligne, colonne) du coin haut gauche d'un nœud."""
    f, a = n.rsplit("!", 1)
    c1, l1, _, _ = range_boundaries(a)
    return f, l1, c1


def unite(x):
    s = x.strip()
    return len(s) <= 4 or (len(s) <= 14 and bool(re.search("[€%/" + chr(0x2192) + "]", s)))


def nombres(texte):
    return [float(re.sub("[" + ESP + "]", "", m).replace(",", ".")) for m in RE_NB.findall(str(texte))]


def cite(v, texte):
    """Vrai si la valeur figure dans le texte (au facteur 100 ou 1000 près)."""
    x = nombre(v)
    return x is not None and any(abs(abs(x) * k - n) <= 0.005 * max(1.0, n) for n in nombres(texte) for k in (1, 100, 1000, 0.001))


def numeros(txt):
    """« 7, 23 et 48 » donne [(7, 7), (23, 23), (48, 48)] ; « 5 à 11 » donne [(5, 11)]."""
    t, out, i = re.findall(r"\d+|à|-|,|\bet\b", txt), [], 0
    while i < len(t):
        if t[i].isdigit():
            if i + 2 < len(t) and t[i + 1] in ("à", "-") and t[i + 2].isdigit():
                out.append((int(t[i]), int(t[i + 2])))
                i += 3
                continue
            out.append((int(t[i]), int(t[i])))
        i += 1
    return out


def critere(c):
    """Prédicat d'un critère de SOMME.SI : nombre, "=x", "<>x", ">=10", jokers * et ?."""
    if nombre(c) is not None:                                               # nombres et dates (numéro de série)
        return lambda v: nombre(v) == nombre(c)
    op, x = re.match(r"^(<=|>=|<>|=|<|>)?(.*)$", "" if c is None else str(c), re.S).groups()
    op = op or "="
    try:
        n = float(x.replace(",", "."))
        f = {"=": n.__eq__, "<>": n.__ne__, "<": n.__gt__, "<=": n.__ge__, ">": n.__lt__, ">=": n.__le__}[op]
        return (lambda v: nombre(v) is None or f(nombre(v))) if op == "<>" else (lambda v: nombre(v) is not None and f(nombre(v)))
    except ValueError:
        pass
    if op in ("=", "<>"):
        rx = re.compile("^" + re.escape(x).replace(r"\*", ".*").replace(r"\?", ".") + "$", re.I | re.S)
        ok = (lambda v: v in (None, "")) if x == "" else (lambda v: isinstance(v, str) and bool(rx.match(v)))
        return ok if op == "=" else (lambda v: not ok(v))
    cmp = {"<": str.__lt__, "<=": str.__le__, ">": str.__gt__, ">=": str.__ge__}[op]
    return lambda v: isinstance(v, str) and cmp(v.lower(), x.lower())


def egal(v, k):
    if isinstance(k, str) and isinstance(v, str):
        if "*" in k or "?" in k:
            return bool(re.match("^" + re.escape(k).replace(r"\*", ".*").replace(r"\?", ".") + "$", v, re.I | re.S))
        return v.lower() == k.lower()
    return nombre(v) is not None and nombre(v) == nombre(k)


def comparer(a, op, b):
    """Comparaison à la manière d'Excel : casse ignorée, nombres < textes < booléens, cellule vide = 0 ou ""."""
    if a is INCONNU or b is INCONNU or any(isinstance(x, str) and x in ERR for x in (a, b)):
        return INCONNU
    a = ("" if isinstance(b, str) else False if isinstance(b, bool) else 0.0) if a is None else a
    b = ("" if isinstance(a, str) else False if isinstance(a, bool) else 0.0) if b is None else b
    rang = lambda x: 2 if isinstance(x, bool) else 1 if isinstance(x, str) else 0
    if rang(a) != rang(b):
        ka, kb = rang(a), rang(b)
    elif isinstance(a, str):
        ka, kb = a.lower(), b.lower()
    else:
        ka, kb = (a, b) if isinstance(a, bool) else (nombre(a), nombre(b))
    if ka is None or kb is None:
        return INCONNU
    return {"=": ka == kb, "<>": ka != kb, "<": ka < kb, ">": ka > kb, "<=": ka <= kb, ">=": ka >= kb}[op]


def vrai(v):
    if v is INCONNU:
        return INCONNU
    if isinstance(v, str):
        return v.upper() == "TRUE" if v.upper() in ("TRUE", "FALSE") else INCONNU
    if v is None or isinstance(v, bool):
        return bool(v)
    return nombre(v) != 0 if nombre(v) is not None else INCONNU


def appels(tokens):
    """Appels imbriqués : (NOM, [argument, ...]) ; un argument est une liste de jetons ou d'appels."""
    res, pile = [], []
    for t in tokens:
        if t.type == Token.FUNC and t.subtype == Token.OPEN:
            nom = t.value[:-1].upper()
            for p in ("_XLFN._XLWS.", "_XLFN.", "_XLWS."):
                nom = nom[len(p):] if nom.startswith(p) else nom
            pile.append((nom, [[]]))
        elif t.type == Token.FUNC and t.subtype == Token.CLOSE and pile:
            ap = pile.pop()
            res.append(ap)
            if pile:
                pile[-1][1][-1].append(ap)
        elif pile and t.type == Token.SEP and t.subtype == Token.ARG:
            pile[-1][1].append([])
        elif pile and t.type != Token.WSPACE:
            pile[-1][1][-1].append(t)
    return res


def tous(items):
    for x in items:
        if isinstance(x, tuple):
            for a in x[1]:
                yield from tous(a)
        else:
            yield x


def jetons(items):
    return {id(t) for t in tous(items)}


def signature(toks, l, c):
    """Formule en notation relative R1C1 : même signature = même formule recopiée."""
    out = []
    for t in toks:
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            p, _, a = t.value.rpartition("!")

            def cv(m):
                ac, lettres, al, num = m.groups()
                cc, ll = CI(lettres), int(num)
                return (f"R{ll}" if al else f"R[{ll - l}]") + (f"C{cc}" if ac else f"C[{cc - c}]")
            out.append((p + "!" if p else "") + RE_A1.sub(cv, a))
        else:
            out.append(t.value)
    return "".join(out)


def enveloppe(chemin):
    """Lecture 0 : ce dont le fichier est fait, lu dans le XML (y compris l'état du cache de chaque formule)."""
    z = zipfile.ZipFile(chemin)
    noms = set(z.namelist())
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    rels = {r.get("Id"): r.get("Target") for r in ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))}
    feuilles, types, dummy, cache, sans = [], Counter(), 0, Counter(), []
    for s in wb.iter(NS + "sheet"):
        cible = rels.get(s.get(RID), "")
        part, t = (cible.lstrip("/") if cible.startswith("/") else "xl/" + cible), Counter()
        if part in noms:
            for _, el in ET.iterparse(z.open(part)):
                if el.tag != NS + "c":
                    continue
                t[el.get("t", "n")] += 1
                fx, v = el.find(NS + "f"), el.find(NS + "v")
                if fx is not None:
                    dummy += bool(fx.text and "DUMMYFUNCTION" in fx.text)
                    k = "sans_valeur" if v is None else "chaine_vide" if el.get("t") == "str" and not v.text else "valeur"
                    cache[k] += 1
                    if k == "sans_valeur" and len(sans) < 20:
                        sans.append(f"{s.get('name')}!{el.get('r')}")
                el.clear()
        feuilles.append({"nom": s.get("name"), "etat": s.get("state", "visible"), "types": dict(t)})
        types.update(t)
    calc = wb.find(NS + "calcPr")
    return {"feuilles": feuilles, "types_cellules": dict(types), "dummyfunction": dummy,
            "formules_cache": dict(cache), "exemples_sans_valeur": sans,
            "noms_definis": len(wb.findall(f".//{NS}definedName")),
            "liens_externes": sum(n.startswith("xl/externalLinks/externalLink") for n in noms),
            "tables": sum(n.startswith("xl/tables/table") for n in noms),
            "tcd": sum(n.startswith("xl/pivotCache/pivotCacheDefinition") for n in noms),
            "power_query": any(n.startswith("customXml/item") and n.endswith(".xml") and b"DataMashup" in z.read(n) for n in noms),
            "macros": "xl/vbaProject.bin" in noms, "calcchain": "xl/calcChain.xml" in noms,
            "calcul": dict(calc.attrib) if calc is not None else {}}


class Classeur:
    def __init__(self, chemin):
        self.f, self.v = load_workbook(chemin), load_workbook(chemin, data_only=True)   # formules, valeurs en cache
        self.noms = {(None, n.upper()): d for n, d in self.f.defined_names.items()}
        for ws in self.f.worksheets:
            self.noms.update({(ws.title, n.upper()): d for n, d in ws.defined_names.items()})
        self.anos, self.sel, self.opaques, self.sig, self.selecteurs = [], [], Counter(), {}, {}
        self.litteraux, self.index, self.durs, self.morts_struct, self.matricielles = [], defaultdict(list), {}, set(), set()
        self._c, self._v, self._m, self._e, self._d, self._h = {}, {}, {}, {}, {}, {}
        self.listes = {}
        for ws in self.f.worksheets:
            for d in ws.data_validations.dataValidation:
                if d.type != "list" or not d.formula1:
                    continue
                f1 = d.formula1.strip()
                if f1.startswith('"'):
                    vals = [x.strip() for x in f1.strip('"').split(",")]
                else:
                    r = self.resoudre(f1.lstrip("="), ws.title)
                    vals = [x for x in plat(self.valeurs(r[0][1], r[0][2])) if x not in (None, "")] if r[0][0] == "plage" else []
                for rg in str(d.sqref).split():
                    c1, l1, c2, l2 = range_boundaries(rg)
                    for n in plat(self.cellules(ws.title, (c1 or 1, l1 or 1, c2 or ws.max_column, min(l2 or ws.max_row, ws.max_row)))):
                        self.listes[n] = vals

    def ano(self, croisement, genre, cellule, detail, sev="à vérifier"):
        self.anos.append({"croisement": croisement, "type": genre, "cellule": cellule, "detail": detail,
                          "severite": sev, "statut": "à qualifier"})

    def dur(self, n, motif):
        self.durs.setdefault(n, motif)

    def resoudre(self, ref, feuille):
        """Liste de ('plage', feuille, (c1, l1, c2, l2)) ou [('opaque', motif)]."""
        r = ref.strip()
        if re.match(r"^'?\[\d+\]", r):
            return [("opaque", "lien externe")]
        if "[" in r:
            return [("opaque", "référence structurée")]
        f, sep, a = r.rpartition("!")
        f, a = (f.strip("'").replace("''", "'") if sep else feuille), a.replace("$", "")
        d = self.noms.get((f, a.upper())) or self.noms.get((None, a.upper()))
        if d is not None:
            try:
                dest = list(d.destinations)
            except Exception:
                dest = []
            return [x for fd, rd in dest for x in self.resoudre(f"'{fd}'!{rd}", fd)] or [("opaque", "nom sans plage")]
        if f not in self.f.sheetnames:
            return [("opaque", "feuille inconnue ou 3D")]
        try:
            c1, l1, c2, l2 = range_boundaries(a)
        except (ValueError, TypeError):
            return [("opaque", "référence illisible")]
        ws = self.f[f]
        return [("plage", f, (c1 or 1, l1 or 1, c2 or ws.max_column, l2 or ws.max_row))]

    def cellules(self, f, b):
        if (f, b) not in self._c:
            self._c[(f, b)] = [[f"{f}!{GL(c)}{l}" for c in range(b[0], b[2] + 1)] for l in range(b[1], b[3] + 1)]
        return self._c[(f, b)]

    def valeurs(self, f, b):
        if (f, b) not in self._v:
            ws = self.v[f]
            self._v[(f, b)] = [[lire(ws, l, c) for c in range(b[0], b[2] + 1)] for l in range(b[1], b[3] + 1)]
        return self._v[(f, b)]

    def valeur(self, n):
        f, l, c = pos(n)
        return lire(self.v[f], l, c)

    def masque(self, f, b, cr):
        if (f, b, repr(cr)) not in self._m:
            p = critere(cr)
            self._m[(f, b, repr(cr))] = frozenset(i for i, v in enumerate(plat(self.valeurs(f, b))) if p(v))
        return self._m[(f, b, repr(cr))]

    def doublons(self, f, b):
        """Clés présentes plusieurs fois dans un vecteur de recherche : {clé: [(position, cellule), ...]}."""
        if (f, b) not in self._d:
            vus = defaultdict(list)
            for i, (n, v) in enumerate(zip(plat(self.cellules(f, b)), plat(self.valeurs(f, b)))):
                if v not in (None, ""):
                    vus[v.lower() if isinstance(v, str) else nombre(v)].append((i, n))
            self._d[(f, b)] = {k: ns for k, ns in vus.items() if len(ns) > 1}
        return self._d[(f, b)]

    def hors(self, f, b, garde):
        """Cellules d'une table qui ne sont ni dans le vecteur de recherche ni dans la colonne renvoyée."""
        if (f, b, garde) not in self._h:
            g = {n for x in garde for n in plat(self.cellules(f, x))}
            self._h[(f, b, garde)] = frozenset(plat(self.cellules(f, b))) - g
        return self._h[(f, b, garde)]

    def arg(self, a, feuille):
        if len(a) == 2 and not any(isinstance(x, tuple) for x in a) and a[0].type == Token.OP_PRE \
                and a[0].value == "-" and a[1].subtype == Token.NUMBER:
            return ("valeur", -float(a[1].value))
        if len(a) != 1:
            return None
        x = a[0]
        if isinstance(x, tuple):                                              # FALSE() écrit par LibreOffice
            return ("valeur", x[0] == "TRUE") if x[0] in ("TRUE", "FALSE") and x[1] in ([[]], []) else ("appel", x)
        if x.type != Token.OPERAND:
            return None
        if x.subtype == Token.RANGE:
            r = self.resoudre(x.value, feuille)
            return ("plage", r[0][1], r[0][2], id(x)) if len(r) == 1 and r[0][0] == "plage" else None
        if x.subtype == Token.NUMBER:
            return ("valeur", float(x.value))
        if x.subtype == Token.TEXT:
            return ("valeur", x.value[1:-1].replace('""', '"'))
        return ("valeur", x.value.upper() == "TRUE") if x.subtype == Token.LOGICAL else None

    def scalaire(self, a, feuille, ici, acc):
        """Valeur d'un argument : constante, cellule unique (valeur en cache) ou EQUIV évaluable."""
        if a is None:
            return None
        if a[0] == "valeur":
            return a[1]
        if a[0] == "plage" and a[2][:2] == a[2][2:]:
            acc["tokens"].add(a[3])
            acc["controle"].add(("plage", a[1], a[2]))
            return lire(self.v[a[1]], a[2][1], a[2][0])
        if a[0] == "appel" and a[1][0] == "MATCH":
            s = self.selection(a[1], feuille, ici)
            if s:
                acc["tokens"] |= s["tokens"]
                acc["controle"] |= s["controle"]
                acc["ecartes"] = s["ecartes"]
                return s["position"]
        return None

    @staticmethod
    def recherche(vec, cle, exact, ordre=1, dernier=False):
        """Position (base 0) renvoyée par une recherche, et indicateur de tri du vecteur."""
        if exact:
            pos_ = [i for i, v in enumerate(vec) if egal(v, cle)]
            return (pos_[-1] if dernier else pos_[0]) if pos_ else None, True
        nums, k = [nombre(v) for v in vec], nombre(cle)
        s = [x for x in nums if x is not None]
        trie = all((a <= b) if ordre == 1 else (a >= b) for a, b in zip(s, s[1:]))
        if k is None or not trie:
            return None, trie
        ok = [i for i, x in enumerate(nums) if x is not None and (x <= k if ordre == 1 else x >= k)]
        return (ok[-1] if ok else None), trie

    def selection(self, ap, f, ici):
        """Masque de lecture d'un appel sélectif dans l'état courant : cellules lues, plages de contrôle, cellules
        que la formule ne peut jamais lire (structurelles). None si l'appel n'est pas évaluable sans moteur."""
        nom, args = ap
        A = [self.arg(a, f) for a in args]
        acc = {"nom": nom, "formule": ici, "tokens": set(), "controle": set(), "effectif": set(),
               "structurel": frozenset(), "boites": [], "position": None, "ecartes": []}
        P = lambda i: A[i] if i < len(A) and A[i] and A[i][0] == "plage" else None
        V = lambda i, defaut=None: self.scalaire(A[i], f, ici, acc) if i < len(A) else defaut
        if nom in SI:
            multi, compte = nom.endswith("S"), nom.startswith("COUNT")
            somme = None if compte else (P(0) if multi else (P(2) or P(0)))
            d = 0 if (compte or not multi) else 1
            paires = [(P(i), V(i + 1)) for i in range(d, len(A) - 1, 2)] if multi else [(P(0), V(1))]
            if not paires or any(p is None or c is None for p, c in paires) or (not compte and somme is None):
                return None
            forme, idx = None, None
            for p, c in paires:
                acc["tokens"].add(p[3])
                acc["controle"].add(("plage", p[1], p[2]))
                h, w = p[2][3] - p[2][1] + 1, p[2][2] - p[2][0] + 1
                if forme not in (None, (h, w)):
                    return None
                forme, m = (h, w), self.masque(p[1], p[2], c)
                idx = m if idx is None else idx & m
            if somme:
                acc["tokens"].add(somme[3])
                (h, w), (c1, l1, c2, l2), cp = forme, somme[2], paires[0][0]
                cel = self.cellules(somme[1], (c1, l1, c1 + w - 1, l1 + h - 1))       # Excel lit la forme du critère
                acc["effectif"] = {cel[i // w][i % w] for i in idx}
                acc["boites"] = [(somme[1], somme[2])]
                if (h, w) != (l2 - l1 + 1, c2 - c1 + 1):
                    acc["structurel"] = frozenset(plat(self.cellules(somme[1], somme[2]))) - frozenset(plat(cel))
                    self.ano("plan×graphe", "plage de somme redimensionnée", ici,
                             f"{nom} : la plage de somme est relue à la forme de la plage de critère", "critique")
                if somme[1] == cp[1] and ((w == 1 and l1 != cp[2][1]) or (h == 1 and c1 != cp[2][0])):
                    self.ano("plan×graphe", "critère et somme décalés", ici,
                             f"{nom} : plage de critère et plage de somme ne commencent pas au même rang", "critique")
            return acc
        if nom == "LOOKUP":
            cle, vec, ret = V(0), P(1), P(2)
            if cle is None or not vec or (vec[2][0] != vec[2][2] and vec[2][1] != vec[2][3]):
                return None
            vals, k = plat(self.valeurs(vec[1], vec[2])), nombre(cle)
            acc["tokens"] |= {vec[3]} | ({ret[3]} if ret else set())
            if k is not None and k >= 1e300:                                  # idiome : dernière valeur numérique
                p = max((i for i, v in enumerate(vals) if nombre(v) is not None), default=None)
            elif isinstance(cle, str) and cle.lower().startswith("zz"):     # idiome : dernier texte
                p = max((i for i, v in enumerate(vals) if isinstance(v, str)), default=None)
            else:
                p, trie = self.recherche(vals, cle, False)
                acc["controle"].add(("plage", vec[1], vec[2]))
                if not trie:
                    self.ano("plan×sens", "recherche approchée sur vecteur non trié", ici, "LOOKUP : résultat imprévisible", "critique")
            src = ret or vec
            cel = plat(self.cellules(src[1], src[2]))
            acc["boites"] = [(src[1], src[2])]
            if p is not None and p < len(cel):
                acc["effectif"] = {cel[p]}
            return acc
        if nom in RECH:
            cle, dernier = V(0), False
            if nom == "XLOOKUP":
                vec, ret, mode, sens = P(1), P(2), V(4, 0), V(5, 1)
                if cle is None or not vec or not ret or mode != 0 or sens not in (1, -1):
                    return None
                fv, vb, exact, ordre, dernier = vec[1], vec[2], True, 1, sens == -1
                retour = plat(self.cellules(ret[1], ret[2]))
                acc["tokens"] |= {vec[3], ret[3]}
                acc["boites"] = [(ret[1], ret[2])]
            else:
                tab = P(1)
                if cle is None or not tab:
                    return None
                acc["tokens"].add(tab[3])
                fv, (c1, l1, c2, l2) = tab[1], tab[2]
                if nom == "MATCH":
                    typ = V(2, 1)
                    if typ is None:
                        return None
                    vb, exact, ordre, retour = tab[2], typ == 0, (-1 if typ == -1 else 1), None
                else:
                    col, approx = V(2), V(3, True)
                    if nombre(col) is None or approx is None:
                        return None
                    col, vert = int(nombre(col)), nom == "VLOOKUP"
                    vb = (c1, l1, c1, l2) if vert else (c1, l1, c2, l1)
                    rb = (c1 + col - 1, l1, c1 + col - 1, l2) if vert else (c1, l1 + col - 1, c2, l1 + col - 1)
                    exact, ordre, retour = approx in (False, 0), 1, plat(self.cellules(fv, rb))
                    acc["boites"] = [(fv, tab[2])]
                    if A[2][0] == "valeur":
                        acc["structurel"] = self.hors(fv, tab[2], (vb, rb))
            acc["controle"].add(("plage", fv, vb))
            p, trie = self.recherche(plat(self.valeurs(fv, vb)), cle, exact, ordre, dernier)
            acc["position"] = None if p is None else p + 1
            if not exact and not trie:
                self.ano("plan×sens", "recherche approchée sur vecteur non trié", ici, f"{nom} : résultat imprévisible", "critique")
            if exact:
                dups = self.doublons(fv, vb)
                acc["ecartes"] = sorted(i for ns in dups.values() for i, _ in (ns[:-1] if dernier else ns[1:]))
                if dups:
                    ex = ", ".join(f"{k!r} ({ns[0][1]}, {ns[1][1]})" for k, ns in list(dups.items())[:3])
                    self.ano("plan×graphe", "clé en doublon", ici,
                             f"{nom} : {len(dups)} clé(s) en doublon, une seule occurrence lue ; ex. {ex}", "critique")
                if retour:
                    acc["structurel"] = acc["structurel"] | {retour[i] for i in acc["ecartes"] if i < len(retour)}
            if retour and p is not None and p < len(retour):
                acc["effectif"] = {retour[p]}
            return acc
        if nom == "INDEX":
            p = P(0)
            if not p or len(A) < 2:
                return None
            acc["ecartes"] = []
            r, er = V(1), acc["ecartes"]
            acc["ecartes"] = []
            c, ec = (V(2) if len(A) > 2 else None), acc["ecartes"]
            if nombre(r) is None or (len(A) > 2 and nombre(c) is None):
                return None
            acc["tokens"].add(p[3])
            cel, acc["boites"] = self.cellules(p[1], p[2]), [(p[1], p[2])]
            r, c = int(nombre(r)), (int(nombre(c)) if c is not None else None)
            if len(cel) == 1 and c is None:
                r, c, er, ec = 1, r, [], er
            c = c or 1
            if 1 <= r <= len(cel) and 1 <= c <= len(cel[0]):
                acc["effectif"] = {cel[r - 1][c - 1]}
            st = {cel[i][c - 1] for i in er if i < len(cel) and c <= len(cel[0])}
            st |= {cel[r - 1][j] for j in ec if 1 <= r <= len(cel) and j < len(cel[0])}
            if all(a and a[0] == "valeur" for a in A[1:3]):
                st |= set(plat(cel)) - acc["effectif"]
            acc["structurel"] = frozenset(st)
            return acc
        if nom == "CHOOSE":
            k = V(0)
            if nombre(k) is None:
                return None
            k, st = int(nombre(k)), set()
            if A[0][0] == "plage" and A[0][2][:2] == A[0][2][2:]:
                for i in range(1, len(A)):
                    self.inscrire(f"{A[0][1]}!{GL(A[0][2][0])}{A[0][2][1]}", float(i), "CHOISIR")
            for i, a in enumerate(A[1:], 1):
                if a and a[0] == "plage":
                    acc["tokens"].add(a[3])
                    acc["boites"].append((a[1], a[2]))
                    cel = plat(self.cellules(a[1], a[2]))
                    if i == k:
                        acc["effectif"] |= set(cel)
                    elif A[0][0] == "valeur":
                        st |= set(cel)
            acc["structurel"] = frozenset(st)
            return acc
        return None

    def inscrire(self, n, valeur, origine):
        """Sélecteur : cellule saisie qui oriente un calcul ; ses valeurs sont celles que les formules testent."""
        s = self.selecteurs.setdefault(n, {"valeurs": [], "origines": set()})
        if valeur is not None and valeur is not INCONNU and not any(comparer(valeur, "=", x) is True for x in s["valeurs"]):
            s["valeurs"].append(valeur)
        s["origines"].add(origine)

    def cellule_saisie(self, x, f):
        """Nom de la cellule si le jeton désigne une seule cellule saisie (pas une formule), sinon None."""
        if isinstance(x, tuple) or x.type != Token.OPERAND or x.subtype != Token.RANGE:
            return None
        r = self.resoudre(x.value, f)
        if len(r) != 1 or r[0][0] != "plage" or r[0][2][:2] != r[0][2][2:]:
            return None
        fb, (c, l) = r[0][1], r[0][2][:2]
        v = lire(self.f[fb], l, c)
        return f"{fb}!{GL(c)}{l}" if v is not None and formule(v) is None else None

    def plage_de(self, arg, f):
        arg = [x for x in arg if isinstance(x, tuple) or x.type != Token.WSPACE]
        if len(arg) != 1 or isinstance(arg[0], tuple) or arg[0].type != Token.OPERAND or arg[0].subtype != Token.RANGE:
            return None
        r = self.resoudre(arg[0].value, f)
        return plat(self.valeurs(r[0][1], r[0][2])) if len(r) == 1 and r[0][0] == "plage" else None

    def evaluer(self, items, f):
        """Valeur en cache d'une expression simple : opérande, comparaison, ET, OU, NON, ESTNUM, NB... ; INCONNU sinon.
        Une cellule saisie comparée à un texte, ou à un nombre non nul par =, est inscrite comme sélecteur."""
        items = [x for x in items if isinstance(x, tuple) or x.type != Token.WSPACE]
        if len(items) == 3 and not isinstance(items[1], tuple) and items[1].type == Token.OP_IN and items[1].value in CMP:
            a, b, op = self.evaluer(items[:1], f), self.evaluer(items[2:], f), items[1].value
            for x, y in ((items[0], items[2]), (items[2], items[0])):
                n = self.cellule_saisie(x, f)
                if n and not isinstance(y, tuple) and y.type == Token.OPERAND and (
                        (y.subtype == Token.TEXT and op in ("=", "<>")) or (y.subtype == Token.NUMBER and op == "=" and float(y.value))):
                    self.inscrire(n, self.evaluer([y], f), "SI")
            return comparer(a, op, b)
        if len(items) == 2 and not isinstance(items[0], tuple) and items[0].type == Token.OP_PRE and items[0].value == "-":
            x = self.evaluer(items[1:], f)
            return -nombre(x) if x is not INCONNU and nombre(x) is not None else INCONNU
        if len(items) != 1:
            return INCONNU
        x = items[0]
        if isinstance(x, tuple):
            nom, args = x
            if nom in ("TRUE", "FALSE"):
                return nom == "TRUE"
            if nom in ("AND", "OR", "NOT"):
                vs = [vrai(self.evaluer(a, f)) for a in args if a]
                if not vs or any(v is INCONNU for v in vs):
                    return INCONNU
                return all(vs) if nom == "AND" else any(vs) if nom == "OR" else not vs[0]
            vals = self.plage_de(args[0], f) if len(args) == 1 else None
            if vals is None:
                return INCONNU
            if nom == "COUNT":
                return float(sum(nombre(v) is not None for v in vals))
            if nom == "COUNTA":
                return float(sum(v not in (None, "") for v in vals))
            if nom in ("SUM", "MIN", "MAX"):
                xs = [nombre(v) for v in vals if nombre(v) is not None]
                return float(sum(xs)) if nom == "SUM" else (min(xs) if nom == "MIN" else max(xs)) if xs else 0.0
            return TESTS[nom](vals[0]) if nom in TESTS and len(vals) == 1 else INCONNU
        if x.type != Token.OPERAND:
            return INCONNU
        if x.subtype == Token.RANGE:
            vals = self.plage_de([x], f)
            return vals[0] if vals is not None and len(vals) == 1 and vals[0] not in ERR else INCONNU
        if x.subtype == Token.NUMBER:
            return float(x.value)
        if x.subtype == Token.TEXT:
            return x.value[1:-1].replace('""', '"')
        return x.value.upper() == "TRUE" if x.subtype == Token.LOGICAL else INCONNU

    def choix(self, ap, f):
        """SI, SI.CONDITIONS, SI.MULTIPLE évalués sur le cache : (jetons des branches mortes, condition sans référence)."""
        nom, args = ap
        vivant, gardes = None, set()
        if nom == "IF" and args:
            c = vrai(self.evaluer(args[0], f))
            if c is INCONNU:
                return None
            vivant, gardes = (1 if c else 2), {0}
        elif nom == "IFS":
            for i in range(0, len(args) - 1, 2):
                c = vrai(self.evaluer(args[i], f))
                if c is INCONNU:
                    return None
                gardes.add(i)
                if c:
                    vivant = i + 1
                    break
        elif nom == "SWITCH" and len(args) >= 3:
            x, gardes = self.evaluer(args[0], f), {0}
            n = self.cellule_saisie(args[0][0], f) if len(args[0]) == 1 else None
            if x is INCONNU:
                return None
            for j in range((len(args) - 1) // 2):
                v = self.evaluer(args[1 + 2 * j], f)
                if v is INCONNU:
                    return None
                gardes.add(1 + 2 * j)
                if n:
                    self.inscrire(n, v, "SI.MULTIPLE")
                if vivant is None and comparer(x, "=", v) is True:
                    vivant = 2 + 2 * j
            if vivant is None and (len(args) - 1) % 2:
                vivant = len(args) - 1
        if vivant is None:
            return None
        morts = set().union(*[jetons(a) for i, a in enumerate(args) if i != vivant and i not in gardes] or [set()])
        fixe = not any(t.type == Token.OPERAND and t.subtype == Token.RANGE for i in gardes for t in tous(args[i]))
        return morts, fixe

    def noeud(self, G, E, ref):
        """Nœud d'une cellule, ou nœud de plage relié à ses cellules non vides."""
        _, f, (c1, l1, c2, l2) = ref
        if (c1, l1) == (c2, l2):
            return f"{f}!{GL(c1)}{l1}"
        k = f"{f}!{GL(c1)}{l1}:{GL(c2)}{l2}"
        if k not in G:
            G.add_node(k, genre="plage")
            E.add_node(k, genre="plage")
            ws = self.f[f]
            if (c2 - c1 + 1) * (l2 - l1 + 1) < len(ws._cells):
                dedans = (ws._cells.get((l, c)) for l in range(l1, l2 + 1) for c in range(c1, c2 + 1))
            else:
                dedans = (x for (l, c), x in ws._cells.items() if c1 <= c <= c2 and l1 <= l <= l2)
            for x in dedans:
                if x is not None and x.value is not None:
                    G.add_edge(f"{f}!{x.coordinate}", k)
                    E.add_edge(f"{f}!{x.coordinate}", k)
        return k

    def graphe(self):
        """G : graphe déclaré (toutes les références). E : graphe effectif (lectures réelles dans l'état courant :
        masques des appels sélectifs, branches mortes des SI, lectures de forme retirées)."""
        G, E = nx.DiGraph(), nx.DiGraph()
        for ws in self.f.worksheets:
            for (l, c), x in sorted(ws._cells.items()):
                if x.value is None:
                    continue
                n, fx = f"{ws.title}!{x.coordinate}", formule(x.value)
                g = "formule" if fx else ("nombre" if nombre(x.value) is not None
                                          and not isinstance(x.value, (datetime, date)) else "texte")
                G.add_node(n, genre=g)
                E.add_node(n, genre=g)
                self.index[(ws.title, l)].append(n)
                if isinstance(x.value, ArrayFormula):
                    self.matricielles |= set(plat(self.cellules(ws.title, range_boundaries(x.value.ref or x.coordinate))))
                if not fx:
                    continue
                try:
                    toks = Tokenizer(fx).items
                except Exception:
                    self.opaques["formule illisible"] += 1
                    continue
                self.sig[(ws.title, l, c)] = signature(toks, l, c)
                aps, morts, fixes, pris = appels(toks), set(), set(), set()
                for ap in aps:
                    if ap[0] in CHOIX:
                        try:
                            r = self.choix(ap, ws.title)
                        except (IndexError, KeyError, TypeError, ValueError, AttributeError):
                            r = None
                        if r:
                            morts |= r[0]
                            fixes |= r[0] if r[1] else set()
                for ap in aps:
                    if ap[0] in OPAQUES:
                        self.opaques[ap[0]] += 1
                    if jetons([ap]) <= morts:
                        continue                                              # appel situé dans une branche morte
                    if ap[0] in FORME:
                        pris |= {id(t) for a in ap[1] for t in a if not isinstance(t, tuple) and t.subtype == Token.RANGE}
                    elif ap[0] in SI | RECH | {"INDEX", "CHOOSE"}:
                        try:
                            s = self.selection(ap, ws.title, n)
                        except (IndexError, KeyError, TypeError, ValueError, AttributeError):
                            s = None
                        if s:
                            self.sel.append(s)
                            pris |= s["tokens"]
                            for m in s["effectif"]:
                                E.add_edge(m, n)
                            for b in s["controle"]:
                                E.add_edge(self.noeud(G, E, b), n)
                for i, t in enumerate(toks):
                    if t.type != Token.OPERAND:
                        continue
                    if t.subtype == Token.NUMBER and id(t) not in morts and any(
                            0 <= j < len(toks) and toks[j].type == Token.OP_IN and toks[j].value in "+-*/^" for j in (i - 1, i + 1)):
                        self.litteraux.append((n, float(t.value)))              # nombre pris dans un calcul, pas un argument
                    if t.subtype != Token.RANGE:
                        continue
                    for r in self.resoudre(t.value, ws.title):
                        if r[0] != "plage":
                            self.opaques[r[1]] += 1
                            continue
                        src = self.noeud(G, E, r)
                        G.add_edge(src, n)
                        if id(t) in fixes:
                            self.morts_struct |= set(plat(self.cellules(r[1], r[2])))
                        if id(t) not in pris and id(t) not in morts:
                            E.add_edge(src, n)
        return G, E

    def segment(self, f, l, c, axe):
        """Signatures des formules contiguës à la cellule, le long d'une ligne ou d'une colonne."""
        d, i, out = ((0, 1) if axe == "ligne" else (1, 0)), 0, []
        while (f, l - (i + 1) * d[0], c - (i + 1) * d[1]) in self.sig:
            i += 1
        x = (l - i * d[0], c - i * d[1])
        while (f, *x) in self.sig:
            out.append(self.sig[(f, *x)])
            x = (x[0] + d[0], x[1] + d[1])
        return out

    def plan(self):
        """Blocs de formules recopiées (signature R1C1 identique et contiguë), ruptures, valeurs écrasées, constantes en bord."""
        parent, sig = {}, self.sig

        def racine(x):
            while parent.setdefault(x, x) != x:
                parent[x] = parent[parent[x]]
                x = parent[x]
            return x
        for (f, l, c), s in sig.items():
            for v in ((f, l, c + 1), (f, l + 1, c)):
                if sig.get(v) == s:
                    parent[racine((f, l, c))] = racine(v)
        groupes = defaultdict(list)
        for k in sig:
            groupes[racine(k)].append(k)
        blocs = [{"feuille": m[0][0], "plage": f"{GL(min(x[2] for x in m))}{min(x[1] for x in m)}:"
                  f"{GL(max(x[2] for x in m))}{max(x[1] for x in m)}", "cellules": len(m), "signature": sig[m[0]]}
                 for m in groupes.values()]
        axes = (("ligne", ((0, -1), (0, 1)), ((-1, 0), (1, 0))), ("colonne", ((-1, 0), (1, 0)), ((0, -1), (0, 1))))
        for (f, l, c), s in sig.items():
            for axe, (a, b), ortho in axes:
                sa = sig.get((f, l + a[0], c + a[1]))
                if not sa or sa != sig.get((f, l + b[0], c + b[1])) or sa == s:
                    continue
                if any(sig.get((f, l + d[0], c + d[1])) == s for d in ortho):
                    continue                                  # ligne ou colonne d'une autre nature (total, ratio)
                lig, tete = self.etiquette(f, l, c)
                lib = tete if axe == "ligne" else lig
                seg = self.segment(f, l, c, axe)
                if RE_SYNTH.search(str(lib or "")) or seg.count(sa) < 2 * seg.count(s) or \
                        (axe == "colonne" and RE_CHK.search(str(lig or ""))) or \
                        self.decale_vide(f, l + a[0], c + a[1], -a[0], -a[1]) or self.aligne(f, l, c):
                    continue
                self.ano("plan×graphe", "rupture de régularité", f"{f}!{GL(c)}{l}",
                         f"formule différente de ses voisines de {axe} ; {'en-tête' if axe == 'ligne' else 'libellé'} {lib!r}", "majeure")
        for ws in self.f.worksheets:
            for (l, c), x in ws._cells.items():
                if formule(x.value) is not None or nombre(x.value) is None or isinstance(x.value, (datetime, date)):
                    continue
                n, lig = f"{ws.title}!{x.coordinate}", self.etiquette(ws.title, l, c)[0]
                for axe, (a, b), ortho in axes:
                    sa = sig.get((ws.title, l + a[0], c + a[1]))
                    if sa and sa == sig.get((ws.title, l + b[0], c + b[1])):
                        if not any(formule(lire(ws, l + d[0], c + d[1])) is None and nombre(lire(ws, l + d[0], c + d[1])) is not None
                                   for d in ortho):                    # sinon : ligne de saisie entre deux lignes de formules
                            self.dur(n, "valeur écrasée")
                            self.ano("plan×graphe", "valeur écrasée", n, f"constante saisie entre deux formules identiques ({axe})", "critique")
                        break
                    for u, v in ((a, b), (b, a)):              # bord : deux formules identiques d'un côté, rien de chiffré de l'autre
                        s1 = sig.get((ws.title, l + v[0], c + v[1]))
                        autre = lire(ws, l + u[0], c + u[1])
                        if s1 and s1 == sig.get((ws.title, l + 2 * v[0], c + 2 * v[1])) and \
                                (autre is None or (isinstance(autre, str) and formule(autre) is None)):
                            self.dur(n, "constante en tête de formules recopiées")
                            self.ano("plan×graphe", "constante en tête de formules recopiées", n,
                                     f"valeur saisie au bord d'une {axe} de formules identiques ; libellé {lig!r}",
                                     "majeure" if lig and RE_CALC.search(lig) else "à vérifier")
        return sorted(blocs, key=lambda b: -b["cellules"])

    def decale_vide(self, f, lv, cv, dl, dc):
        """La formule voisine, recopiée sur la cellule, lirait une cellule vide : l'écart tient à une donnée absente."""
        for t in Tokenizer(formule(lire(self.f[f], lv, cv))).items:
            if t.type == Token.OPERAND and t.subtype == Token.RANGE and ":" not in t.value:
                p, _, a = t.value.rpartition("!")
                m, fb = re.fullmatch(r"(\$?)([A-Z]{1,3})(\$?)(\d+)", a), (p.strip("'").replace("''", "'") if p else f)
                if m and fb in self.f.sheetnames:
                    c, l = CI(m.group(2)) + (0 if m.group(1) else dc), int(m.group(4)) + (0 if m.group(3) else dl)
                    if l >= 1 and c >= 1 and lire(self.f[fb], l, c) is None:
                        return True
        return False

    def aligne(self, f, l, c):
        """La formule lit, dans une autre feuille, une ligne qui porte le même libellé qu'elle : l'écart de position est voulu."""
        norme = lambda x: re.sub(r"\s+", " ", str(x or "")).strip().lower()
        lig = norme(self.etiquette(f, l, c)[0])
        for t in Tokenizer(formule(lire(self.f[f], l, c))).items if lig else []:
            if t.type == Token.OPERAND and t.subtype == Token.RANGE and "!" in t.value:
                for r in self.resoudre(t.value, f):
                    if r[0] == "plage" and r[1] != f and norme(self.etiquette(r[1], r[2][1], r[2][0])[0]) == lig:
                        return True
        return False

    def etiquette(self, f, l, c):
        """Libellé de ligne (premier texte à gauche, unités écartées), en-tête de colonne (texte, date ou année au-dessus)."""
        if (f, l, c) not in self._e:
            ws, lig, col, un = self.v[f], None, None, None
            an = lambda x: isinstance(x, int) and 1950 <= x <= 2149
            for j in range(c - 1, 0, -1):
                x = lire(ws, l, j)
                if isinstance(x, str) and x.strip() and x not in ERR:
                    if unite(x) or f"{f}!{GL(j)}{l}" in self.selecteurs or f"{f}!{GL(j)}{l}" in self.listes:
                        un = un or x
                        continue
                    lig = x
                    break
            for i in range(l - 1, 0, -1):
                x = lire(ws, i, c)
                if isinstance(x, (datetime, date)) or (isinstance(x, str) and x not in ERR) or \
                        (an(x) and (an(lire(ws, i, c - 1)) or an(lire(ws, i, c + 1)))):
                    col = x.isoformat()[:10] if isinstance(x, (datetime, date)) else str(x)
                    break
            self._e[(f, l, c)] = (lig or un, col)
        return self._e[(f, l, c)]

    def libelles(self):
        """Dictionnaire positionné de tout ce qui porte du sens (textes saisis ou calculés, dates, noms, tables,
        commentaires, listes de validation, noms de feuilles)."""
        out = []
        for ws in self.f.worksheets:
            wv = self.v[ws.title]
            out.append({"cellule": ws.title, "texte": ws.title, "origine": "feuille"})
            for (l, c), x in ws._cells.items():
                if x.value is None:
                    continue
                fx, adr = formule(x.value), f"{ws.title}!{x.coordinate}"
                val = lire(wv, l, c) if fx else x.value
                if isinstance(val, str) and val and val not in ERR:
                    out.append({"cellule": adr, "texte": val, "origine": "formule" if fx else "saisi"})
                elif isinstance(val, (datetime, date)):
                    out.append({"cellule": adr, "texte": val.isoformat()[:10], "origine": "date"})
                if x.comment:
                    out.append({"cellule": adr, "texte": x.comment.text, "origine": "commentaire"})
            out += [{"cellule": f"{ws.title}!{t.ref}", "texte": f"{t.name}[{k.name}]", "origine": "table"}
                    for t in ws.tables.values() for k in t.tableColumns]
            out += [{"cellule": f"{ws.title}!{d.sqref}", "texte": d.formula1, "origine": "liste"}
                    for d in ws.data_validations.dataValidation if d.type == "list"]
        out += [{"cellule": d.attr_text, "texte": nom, "origine": "nom" + (f" ({p})" if p else "")}
                for (p, nom), d in self.noms.items()]
        return out

    def periodes(self, n):
        f, l, c = pos(n)
        lig, tete = self.etiquette(f, l, c)
        return set(RE_AN.findall(f"{lig or ''} {tete or ''}")), lig

    def sorties(self, specs):
        if not specs and "Map" in self.f.sheetnames:
            rows = list(self.v["Map"].iter_rows(values_only=True))
            h = [str(x).lower() if x else "" for x in rows[0]] if rows else []
            if {"role", "feuille", "plage"} <= set(h):
                specs = [f"'{r[h.index('feuille')]}'!{r[h.index('plage')]}" for r in rows[1:] if r[h.index("role")] == "output"]
        cel = []
        for s in specs:
            cel += [n for r in self.resoudre(s, self.f.sheetnames[0]) if r[0] == "plage" for n in plat(self.cellules(r[1], r[2]))]
        return cel

    def constantes(self, G, cone=frozenset()):
        """Valeur saisie là où le libellé annonce un calcul ; nombre écrit dans une formule qui recopie une hypothèse."""
        for ws in self.f.worksheets:
            if RE_DOC.search(ws.title):
                continue
            for (l, c), x in ws._cells.items():
                v = x.value
                if formule(v) is not None or nombre(v) is None or isinstance(v, (datetime, date)):
                    continue
                lig = self.etiquette(ws.title, l, c)[0]
                if lig and RE_CALC.search(lig) and not cite(v, lig):
                    n = f"{ws.title}!{x.coordinate}"
                    self.dur(n, "calcul annoncé par le libellé")
                    self.ano("plan×sens", "valeur saisie là où le libellé annonce un calcul", n,
                             f"libellé {lig!r} ; valeur saisie {v:g}", "majeure")
        entrees = defaultdict(list)
        for n in G:
            if G.nodes[n].get("genre") == "nombre" and G.out_degree(n):
                v = nombre(self.valeur(n))
                if v is not None and abs(v) not in TRIVIAUX and not (v == int(v) and 1950 <= v <= 2149):
                    entrees[round(v, 9)].append(n)
        for n, v in self.litteraux:
            cands = entrees.get(round(v, 9), [])
            if abs(v) in TRIVIAUX or not cands:
                continue
            meme = [x for x in cands if pos(x)[0] == pos(n)[0]]
            if set(cands) & set(G.predecessors(n)) or len(meme) > 3 or (not meme and len(cands) > 1):
                continue
            self.ano("graphe×valeurs", "hypothèse recopiée en dur", n,
                     f"le nombre {v:g} écrit dans la formule vaut {', '.join(meme or cands)} ; la formule ne suivra pas cette hypothèse",
                     ("critique" if n in cone else "majeure") if meme else "à vérifier")

    def alias(self):
        """Noms sous lesquels une feuille est citée : nom complet, nom sans préfixe, sigle en capitales (CR)."""
        out = []
        for s in self.f.sheetnames:
            out.append((s, s, False))
            court = re.sub(r"^(att\.|détail|detail|input|hyp\.|synthèse|synthese)\s+", "", s, flags=re.I)
            if court != s and len(court) > 3:
                out.append((court, s, False))
            mots = [w for w in re.findall(r"[^\W\d_]+", s) if w.lower() not in MOTS_VIDES]
            if len(mots) >= 2:
                out.append(("".join(w[0] for w in mots).upper(), s, True))
        return sorted(out, key=lambda x: -len(x[0]))

    @staticmethod
    def occurrences(texte, al):
        prises = []
        for a, s, casse in al:
            for m in re.finditer(r"(?<!\w)" + re.escape(a) + r"(?!\w)", texte, 0 if casse else re.I):
                if not any(m.start() < e and d < m.end() for d, e, _ in prises):
                    prises.append((m.start(), m.end(), s))
        return sorted(prises)

    def liens(self, G, f, lignes, cible, amont_seul=False):
        """Lignes de la feuille cible reliées aux lignes données de la feuille f : lues par elles, ou qui les lisent."""
        out = set()
        for l in lignes:
            for n in self.index.get((f, l), ()):
                for m in ([] if amont_seul else G.successors(n)):
                    for y in (G.successors(m) if G.nodes[m].get("genre") == "plage" else (m,)):
                        if pos(y)[0] == cible:
                            out.add(pos(y)[1])
                for m in G.predecessors(n):
                    fm, a = m.rsplit("!", 1)
                    if fm == cible:
                        _, l1, _, l2 = range_boundaries(a)
                        out.update(range(l1, l2 + 1))
        return out

    def titre(self, f, l):
        """Libellé d'une ligne : premier texte porteur de sens en partant de la gauche."""
        for c in range(1, min(self.f[f].max_column, 12) + 1):
            x = lire(self.v[f], l, c)
            if isinstance(x, str) and len(x.strip()) > 1 and x not in ERR and not unite(x):
                return x.strip()[:40]
        return ""

    def renvois(self, G):
        """« ligne N » cité dans un texte : la ligne citée doit être celle que les formules relient réellement."""
        al = self.alias()
        for ws in self.f.worksheets:
            doc = bool(RE_DOC.search(ws.title))
            for (l, c), x in list(ws._cells.items()):
                t = x.value
                if not isinstance(t, str) or formule(t) or not RE_LIGNE.search(t):
                    continue
                occ, sujet = self.occurrences(t, al), ws.title
                if doc:                                      # page de documentation : le sujet est l'onglet nommé à gauche
                    sujet = next((o[0][2] for o in (self.occurrences(str(lire(ws, l, j) or ""), al) for j in range(1, c)) if o), None)
                    if not sujet:
                        continue
                decrit = doc or not any(G.nodes[v].get("genre") == "formule" for v in self.index.get((ws.title, l), ()))
                lignes = sorted({k[1] for k in self.index if k[0] == sujet}) if decrit else [l]
                for m in RE_LIGNE.finditer(t):
                    cible = next((o[2] for o in occ if 0 <= m.start() - o[1] <= 3 or (
                        0 <= o[0] - m.end() <= 8 and re.match(r"\s*(du|de la|de l'|de|des)\s", t[m.end():]))), None)
                    cf = cible or (sujet if decrit else ws.title)
                    nums = [(a, b) for a, b in numeros(m.group(1)) if 1 <= a <= b <= self.f[cf].max_row + 100]
                    if not nums:
                        continue
                    n, txt = f"{ws.title}!{x.coordinate}", m.group(0)
                    if decrit and cf == sujet:                # description d'un onglet : la ligne citée doit exister
                        vides = [(a, b) for a, b in nums if not self.index.get((cf, a)) or not self.index.get((cf, b))]
                        if vides:
                            self.ano("graphe×sens", "renvoi vers une ligne vide", n, f"« {txt} » de {cf} : ligne vide ou absente", "majeure")
                        continue
                    R = self.liens(G, sujet, lignes, cf) - ({l} if cf == ws.title and not decrit else set())
                    if R and not decrit and cf != sujet:        # la ligne peut se composer d'autres lignes de sa feuille
                        R |= self.liens(G, sujet, self.liens(G, sujet, [l], sujet, True) - {l}, cf)
                    mots = lambda x: {w.lower() for w in re.findall(r"[^\W\d_]{4,}", x)} - {"ligne", "lignes"}
                    if not R:                                  # simple renvoi de lecture : admis si la ligne citée porte le même mot
                        if cible and cible != ws.title and not decrit and not any(
                                mots(t) - mots(cf) & mots(self.titre(cf, i)) for a, b in nums for i in range(a, b + 1)):
                            self.ano("graphe×sens", "renvoi sans liaison", n,
                                     f"« {txt} » de {cf} : aucune formule ne relie cette ligne à {cf}", "à vérifier")
                        continue
                    faux = [(a, b) for a, b in nums if not any(a <= r <= b for r in R)]
                    if faux:
                        vrai_ = ", ".join(f"{r} « {self.titre(cf, r)} »" for r in sorted(R)[:4])
                        self.ano("graphe×sens", "renvoi de ligne faux", n,
                                 f"« {txt} » de {cf} (ligne {faux[0][0]} : « {self.titre(cf, faux[0][0])} ») ; les formules relient : {vrai_}", "majeure")

    def domaines(self, G):
        """Sélecteurs et leurs valeurs possibles : liste de validation, cas testés par SI ou SI.MULTIPLE, rang de CHOISIR."""
        for n, vals in self.listes.items():
            f, l, c = pos(n)
            if n in G and G.out_degree(n) and formule(lire(self.f[f], l, c)) is None:
                self.selecteurs.setdefault(n, {"valeurs": [], "origines": set()})["liste"] = vals
        out = {}
        for n, s in self.selecteurs.items():
            cur, liste = self.valeur(n), s.get("liste")
            vals = list(liste or [])
            for v in s["valeurs"] + [cur]:
                if v is not None and not any(comparer(v, "=", x) is True for x in vals):
                    vals.append(v)
            out[n] = {"valeurs": vals, "cas": list(liste or s["valeurs"]), "courante": cur,
                      "origines": sorted(s["origines"] | ({"liste"} if liste else set())), "formules": len(list(suivants(G, n)))}
            if liste and not any(comparer(cur, "=", x) is True for x in liste):
                self.ano("plan×sens", "sélecteur hors de sa liste", n, f"valeur {cur!r} absente de la liste {liste}", "critique")
            elif not liste and s["origines"] & {"SI", "SI.MULTIPLE"} and not any(comparer(cur, "=", x) is True for x in s["valeurs"]):
                self.ano("graphe×sens", "sélecteur hors des cas prévus", n,
                         f"valeur {cur!r} ; cas testés : {'; '.join(map(str, s['valeurs']))} ; la formule prend sa branche par défaut", "à vérifier")
        return out


def suivants(G, n):
    """Formules qui lisent n, directement ou par un nœud de plage."""
    for m in G.successors(n):
        yield from (G.successors(m) if G.nodes[m].get("genre") == "plage" else (m,))


def regrouper(anos, cone=frozenset()):
    """Une ligne par (croisement, type, détail) : les formules recopiées ne gonflent pas le registre."""
    g = {}
    for a in anos:
        e = g.setdefault((a["croisement"], a["type"], a["detail"]), dict(a, cellules=[]))
        if a["cellule"] not in e["cellules"]:
            e["cellules"].append(a["cellule"])
    for e in g.values():
        cs = e.pop("cellules")
        e["dans_cone"] = any(c in cone for c in cs)
        e["cellule"], e["occurrences"] = ", ".join(cs[:6]) + (f" (+{len(cs) - 6})" if len(cs) > 6 else ""), len(cs)
    return sorted(g.values(), key=lambda x: (SEV.index(x["severite"]), not x["dans_cone"], -x["occurrences"]))


def proche(a, b, tol=1e-6):
    if isinstance(a, str) and a in ERR and isinstance(b, str) and b.startswith("#"):
        return True
    if a in (None, "") and b in (None, ""):
        return True
    u, v = nombre(a), nombre(b)
    if u is not None and v is not None:
        return abs(u - v) <= tol * max(1.0, abs(u), abs(v))
    return a == b


class Moteur:
    """LibreOffice piloté par UNO sur une copie : le classeur reste chargé et chaque essai ne recalcule que ce qui
    dépend de la cellule modifiée (quelques millisecondes). L'original n'est jamais ouvert en écriture."""

    def __init__(self, chemin, delai=90):
        import uno
        from com.sun.star.beans import PropertyValue
        self.tmp = Path(tempfile.mkdtemp())
        copie = self.tmp / ("copie" + Path(chemin).suffix)
        shutil.copy(chemin, copie)
        nom = f"af{os.getpid()}_{int(time.time() * 1000) % 100000}"
        self.proc = subprocess.Popen(["soffice", "--headless", "--invisible", "--nologo", "--norestore", "--nodefault",
                                      "--nolockcheck", f"-env:UserInstallation={(self.tmp / 'profil').as_uri()}",
                                      f"--accept=pipe,name={nom};urp;"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        ctx = uno.getComponentContext()
        res = ctx.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", ctx)
        fin = time.time() + delai
        while True:
            try:
                rctx = res.resolve(f"uno:pipe,name={nom};urp;StarOffice.ComponentContext")
                break
            except Exception:
                if time.time() > fin or self.proc.poll() is not None:
                    self.fermer()
                    raise RuntimeError("LibreOffice ne répond pas")
                time.sleep(0.3)
        bureau = rctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", rctx)
        p = PropertyValue()
        p.Name, p.Value = "Hidden", True
        self.doc = bureau.loadComponentFromURL(copie.as_uri(), "_blank", 0, (p,))
        self.doc.calculateAll()                                   # base : le calcul de LibreOffice, pas le cache
        self._c = {}

    def cellule(self, n):
        if n not in self._c:
            f, a = n.rsplit("!", 1)
            self._c[n] = self.doc.Sheets.getByName(f).getCellRangeByName(a)
        return self._c[n]

    def valeur(self, n):
        c = self.cellule(n)
        if c.getError():
            return f"#ERR{c.getError()}"
        t = c.getType().value
        if t == "EMPTY":
            return None
        return c.getString() if t == "TEXT" or (t == "FORMULA" and c.FormulaResultType2 == 2) else c.getValue()

    def etat(self, n):
        c = self.cellule(n)
        t = c.getType().value
        return ("f", c.getFormula()) if t == "FORMULA" else None if t == "EMPTY" else c.getString() if t == "TEXT" else c.getValue()

    def poser(self, n, v):
        c = self.cellule(n)
        if isinstance(v, tuple):
            c.setFormula(v[1])
        elif v is None:
            c.setFormula("")
        elif isinstance(v, str):
            c.setString(v)
        else:
            c.setValue(float(v))

    def sorties(self, S):
        return [self.valeur(s) for s in S]

    def enregistrer(self, chemin):
        """Copie recalculée au format xlsx : les valeurs en cache y sont celles de LibreOffice."""
        from com.sun.star.beans import PropertyValue
        p = PropertyValue()
        p.Name, p.Value = "FilterName", "Calc MS Excel 2007 XML"
        self.doc.storeToURL(Path(chemin).resolve().as_uri(), (p,))

    def fermer(self):
        try:
            self.doc.close(True)
        except Exception:
            pass
        try:
            self.proc.terminate()
            self.proc.wait(15)
        except Exception:
            self.proc.kill()
        shutil.rmtree(self.tmp, ignore_errors=True)


def ecarts(o, base, tol=1e-9):
    """Écart de chaque sortie : nombre, 0.0 si inchangée, 'modifiée' pour un texte ou une erreur."""
    out = []
    for a, b in zip(o, base):
        u, v = nombre(a), nombre(b)
        if u is not None and v is not None:
            out.append(u - v if abs(u - v) > tol * max(1.0, abs(v)) else 0.0)
        else:
            out.append(0.0 if a == b else "modifiée")
    return out


def bouge(d):
    return any(x != 0 for x in d)


def variantes(v):
    """Essais successifs : faible (+1 %) pour l'effet chiffré, forts (×2, ×0) pour les seuils éloignés."""
    if isinstance(v, str):
        return [("texte modifié", v + " ¤")]
    if v:
        return [("+1 %", v * 1.01), ("×2", v * 2), ("×0", 0.0)]
    return [("+1", 1.0), ("+1000", 1000.0), ("-1000", -1000.0)]


def mesurer(m, cellules, sorties, tol=1e-9):
    """Effet de chaque cellule sur chaque sortie, un essai à la fois, valeur restaurée après chaque essai."""
    base, out = m.sorties(sorties), {}
    for n in cellules:
        e, v = m.etat(n), m.valeur(n)
        if e is None or (isinstance(v, str) and v.startswith("#ERR")):
            continue
        r = {"valeur": v, "classe": "sans effet"}
        for i, (lab, w) in enumerate(variantes(v)):
            try:
                m.poser(n, w)
                d = ecarts(m.sorties(sorties), base, tol)
            except Exception:                                  # feuille protégée, cellule fusionnée
                r["classe"] = "non testée"
                break
            finally:
                try:
                    m.poser(n, e)
                except Exception:
                    pass
            if i == 0:
                r["delta"], r["essai"] = d, lab
            if bouge(d):
                r["classe"] = "pesante" if i == 0 else "seuil"
                if i:
                    r["seuil"] = {"essai": lab, "delta": d}
                break
        out[n] = r
    return out, base, ecarts(m.sorties(sorties), base, tol)


def absorbeurs(m, G, cone, n, base_vals, tol=1e-9, limite=500):
    """Formules où meurt la perturbation d'une cellule sans effet : elles lisent une cellule qui a bougé sans bouger."""
    e, v = m.etat(n), m.valeur(n)
    m.poser(n, variantes(v)[1][1] if len(variantes(v)) > 1 else variantes(v)[0][1])
    vus, arret, file = {n}, set(), [n]
    try:
        while file and len(vus) < limite:
            x = file.pop()
            for y in suivants(G, x):
                if y in vus or y not in cone:
                    continue
                vus.add(y)
                (arret.add if proche(m.valeur(y), base_vals.get(y), tol) else file.append)(y)
    finally:
        m.poser(n, e)
    return sorted(arret)


def etats(m, sels, sorties, base, sans_effet, tol=1e-9, maxi=12, budget=40000):
    """Chaque sélecteur essayé sur chacune de ses valeurs ; pour un sélecteur global, les cellules sans effet sont
    réessayées dans cet état (celles qui s'y réveillent sont conjoncturelles)."""
    out, essais = [], 0
    for n, info in sorted(sels.items(), key=lambda x: -x[1]["formules"]):
        e, cur = m.etat(n), m.valeur(n)
        if e is None:
            continue
        vals = [v for v in info["valeurs"] if comparer(v, "=", cur) is not True][:maxi] or ["¤ autre valeur"]
        for v in vals:
            w = v if isinstance(v, str) else float(v) if isinstance(v, bool) else nombre(v)
            if w is None:
                continue
            m.poser(n, w)
            o, rev = m.sorties(sorties), []
            if info["formules"] >= 3 and essais + 3 * len(sans_effet) <= budget:
                essais += 3 * len(sans_effet)
                for x in sans_effet:
                    ex, vx = m.etat(x), m.valeur(x)
                    for _, w in variantes(vx):
                        m.poser(x, w)
                        d = ecarts(m.sorties(sorties), o, tol)
                        m.poser(x, ex)
                        if bouge(d):
                            rev.append(x)
                            break
            m.poser(n, e)
            out.append({"selecteur": n, "valeur": v, "ecart": ecarts(o, base, tol), "reveillees": rev})
    return out


def topologie(G, sorties, cone, poids=None):
    """Points de passage obligés (dominateurs vers la sortie), profondeur et nombre de chemins, sur le cône déclaré."""
    H, R0 = G.subgraph(cone).copy(), "⊤"
    H.add_edges_from((s, R0) for s in sorties if s in H)
    idom = nx.immediate_dominators(H.reverse(copy=False), R0)
    feuilles = [n for n in H if n != R0 and H.in_degree(n) == 0 and H.nodes[n].get("genre") != "plage"]
    compte, masse = Counter(), Counter()
    total = sum(poids.get(x, 0) for x in feuilles) if poids else 0
    for x in feuilles:
        y = idom.get(x)
        while y is not None and y != R0:
            compte[y] += 1
            masse[y] += poids.get(x, 0) if poids else 0
            y = idom.get(y)
    res = {"hypotheses": len(feuilles), "cycles": not nx.is_directed_acyclic_graph(H), "_prof": {}, "_chemins": {}}
    res["points_obliges"] = [{"cellule": d, "hypotheses": k, "part": round(k / max(1, len(feuilles)), 3),
                              "part_effet": round(masse[d] / total, 3) if total else None}
                             for d, k in compte.most_common() if H.nodes[d].get("genre") == "formule" and d not in sorties][:12]
    if res["cycles"]:
        return res
    prof, chem = {R0: 0}, {R0: 1}
    for x in reversed(list(nx.topological_sort(H))):
        if x != R0:
            s = list(H.successors(x))
            prof[x] = max((prof[y] for y in s), default=0) + (0 if H.nodes[x].get("genre") == "plage" else 1)
            chem[x] = sum(chem[y] for y in s)
    cellules = [x for x in prof if x != R0 and H.nodes[x].get("genre") != "plage"]
    x, chaine = max(cellules, key=lambda n: prof[n], default=None), []
    while x is not None and x != R0:
        if H.nodes[x].get("genre") != "plage":
            chaine.append(x)
        x = max(H.successors(x), key=lambda y: prof[y], default=None)
    res.update({"_prof": prof, "_chemins": chem, "profondeur_max": max((prof[c] for c in cellules), default=0) - 1,
                "chaine": chaine, "chemins_max": [{"cellule": c, "chemins": chem[c]} for c in
                                                  sorted(feuilles, key=lambda c: -chem[c])[:8]]})
    return res


def dessiner(k, G, cone, sorties, env, topo, poids, dossier):
    """Carte d'ensemble (flux entre feuilles, par niveaux) et carte des couloirs (une bande par feuille, de la saisie à
    la sortie, points de passage obligés et chaîne la plus longue)."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Patch
    logging.getLogger("matplotlib.font_manager").setLevel(logging.ERROR)
    plt.rcParams.update({"font.family": "sans-serif", "font.sans-serif": ["Arial", "Liberation Sans", "DejaVu Sans"], "font.size": 8})
    d = Path(dossier)
    d.mkdir(parents=True, exist_ok=True)
    fe = lambda n: n.rsplit("!", 1)[0]
    noms = [f["nom"] for f in env["feuilles"] if f["nom"] in k.f.sheetnames]
    masquees = {f["nom"] for f in env["feuilles"] if f["etat"] != "visible"}
    W = Counter((fe(u), fe(v)) for u, v in G.edges() if fe(u) != fe(v))
    Wc = {e for u, v in G.edges() if u in cone and v in cone and (e := (fe(u), fe(v)))[0] != e[1]}
    S, Sc = nx.DiGraph(), nx.DiGraph()
    S.add_nodes_from(noms)
    S.add_edges_from(W)
    Sc.add_edges_from(Wc)                                         # niveaux fixés par le seul flux utile à la sortie
    C, niv, couche = nx.condensation(Sc), {}, {}
    for x in nx.topological_sort(C):
        niv[x] = max((niv[p] + 1 for p in C.predecessors(x)), default=0)
        couche.update({n: niv[x] for n in C.nodes[x]["members"]})
    for _ in range(len(noms)):                                    # îlots : après ce qu'ils lisent, sinon avant ce qu'ils nourrissent
        for n in noms:
            if n not in couche and S.degree(n):
                pr = [couche[p] for p in S.predecessors(n) if p in couche]
                su = [couche[q] for q in S.successors(n) if q in couche]
                if pr or su:
                    couche[n] = max(pr) + 1 if pr else max(0, min(su) - 1)
    col, y = defaultdict(list), {}
    dans = Counter(fe(n) for n in cone if G.nodes[n].get("genre") != "plage")
    for n in noms:
        if S.degree(n):
            col[couche.get(n, 0)].append(n)
    for c in col:
        y.update({n: i for i, n in enumerate(col[c])})
    for _ in range(2):                    # feuilles du cône en haut de chaque colonne, puis ordre par barycentre des prédécesseurs
        for c in sorted(col):
            col[c].sort(key=lambda n: (not dans[n], sum(y[p] for p in S.predecessors(n)) / S.in_degree(n) if S.in_degree(n) else y[n]))
            y.update({n: i for i, n in enumerate(col[c])})
    haut, P = max([len(v) for v in col.values()] + [1]), {}
    for c, ns in col.items():
        for i, n in enumerate(ns):
            P[n] = (c * 3.6, -(i + (haut - len(ns)) / 2) * 1.2)
    isoles = [n for n in noms if not S.degree(n)]
    for i, n in enumerate(isoles):
        P[n] = ((i % max(1, len(col))) * 3.6, -(haut + 0.6 + i // max(1, len(col))) * 1.2)
    nform, fsort = Counter(f for f, _, _ in k.sig), {fe(s) for s in sorties}
    larg = max(x for x, _ in P.values()) + 3.6
    bas = min(yy for _, yy in P.values()) - 1
    fig, ax = plt.subplots(figsize=(max(7.0, larg * 0.95), max(3.5, -bas * 0.8 + 1.2)))
    for (a, b), w in W.items():
        (x1, y1), (x2, y2) = P[a], P[b]
        utile = (a, b) in Wc
        coul = "#1F4E79" if utile else "#BFBFBF"
        if x2 > x1:
            dep, arr, rad = (x1 + 1.3, y1), (x2 - 1.3, y2), 0.0
        elif x2 < x1:
            dep, arr, rad = (x1 - 1.3, y1), (x2 + 1.3, y2), 0.3
        else:                                                     # même colonne : arc par la droite
            dep, arr, rad = (x1 + 1.3, y1 - 0.15), (x2 + 1.3, y2 + 0.15), (-0.5 if y2 < y1 else 0.5)
        ax.add_patch(FancyArrowPatch(dep, arr, arrowstyle="-|>", mutation_scale=9, lw=0.5 + 1.1 * math.log10(1 + w), color=coul,
                                     alpha=0.85 if utile else 0.6, connectionstyle=f"arc3,rad={rad}", zorder=1 if not utile else 2))
        if utile and x2 != x1:
            ax.text(dep[0] + 0.15 * (arr[0] - dep[0]), dep[1] + 0.15 * (arr[1] - dep[1]), str(w), fontsize=6, color=coul, ha="center",
                    va="center", bbox=dict(fc="white", ec="none", pad=0.3), zorder=4)
    for n, (x, yy) in P.items():
        fc, ec, tc = ("#1F4E79", "#1F4E79", "white") if n in fsort else ("#DCE6F2", "#1F4E79", "#1F1F1F") if dans[n] \
            else ("#F2F2F2", "#A6A6A6", "#7F7F7F")
        ax.add_patch(FancyBboxPatch((x - 1.3, yy - 0.38), 2.6, 0.76, boxstyle="round,pad=0.02,rounding_size=0.08", fc=fc, ec=ec,
                                    lw=1.1, ls="--" if n in masquees else "-", zorder=3))
        ax.text(x, yy + 0.13, n if len(n) <= 26 else n[:25] + "…", ha="center", va="center", fontsize=8, weight="bold", color=tc, zorder=4)
        ax.text(x, yy - 0.17, f"{nform[n]} formules · {dans[n]} dans le cône", ha="center", va="center", fontsize=6.3, color=tc, zorder=4)
    if isoles:
        ax.text(-1.3, -(haut + 0.1) * 1.2, "Sans lien avec les autres feuilles", fontsize=7, color="#7F7F7F", va="center")
    ax.set_xlim(-1.6, larg - 2.0)
    ax.set_ylim(bas, 0.9)
    ax.axis("off")
    ax.set_title(f"Carte d'ensemble : flux entre feuilles vers {', '.join(sorties[:2])}{'…' if len(sorties) > 2 else ''} "
                 "(bleu : flux utile à la sortie, avec le nombre de liens ; gris : autres liens)", fontsize=9.5, loc="left")
    ax.legend(handles=[Patch(fc="#1F4E79", label="Feuille de sortie"), Patch(fc="#DCE6F2", ec="#1F4E79", label="Dans le cône"),
                       Patch(fc="#F2F2F2", ec="#A6A6A6", label="Îlot, hors cône"),
                       Patch(fc="white", ec="#7F7F7F", ls="--", label="Feuille masquée")],
              loc="lower right", fontsize=6.5, frameon=False, ncol=4)
    f1 = d / "carte_ensemble.png"
    fig.savefig(f1, dpi=160, bbox_inches="tight")
    plt.close(fig)
    prof = topo.get("_prof") or {}
    if not prof:
        return [str(f1)]
    dist = {n: p - 1 for n, p in prof.items() if n != "⊤" and G.nodes[n].get("genre") != "plage"}
    bandes = sorted({fe(n) for n in dist}, key=lambda f: (couche.get(f, 0), noms.index(f) if f in noms else 0))
    L, dmax, grp = {f: i for i, f in enumerate(bandes)}, max(dist.values()), defaultdict(list)
    for n, dd in dist.items():
        grp[(fe(n), dd)].append(n)
    xy = lambda g: (dmax - g[1], -L[g[0]])
    A = Counter()
    for u, v in G.subgraph(cone).edges():
        for p in ([u] if G.nodes[u].get("genre") != "plage" else G.predecessors(u)) if G.nodes[v].get("genre") != "plage" else []:
            if p in dist and v in dist:
                A[((fe(p), dist[p]), (fe(v), dist[v]))] += 1
    tot = max((sum(poids.get(n, 0) for n in ns) for ns in grp.values()), default=0) if poids else 0
    fig, ax = plt.subplots(figsize=(max(8.0, 0.6 * (dmax + 1) + 3), max(3.0, 0.62 * len(bandes) + 1.6)))
    for i in range(len(bandes)):
        ax.axhspan(-i - 0.5, -i + 0.5, color="#F4F7FB" if i % 2 else "white", zorder=0)
        ax.text(-0.7, -i, bandes[i], ha="right", va="center", fontsize=7.5)
    for (g1, g2), w in A.items():
        (x1, y1), (x2, y2) = xy(g1), xy(g2)
        ax.plot([x1, x2], [y1, y2], color="#8C8C8C", lw=0.25 + 0.5 * math.log10(1 + w), alpha=0.35, zorder=1)
    cmap = plt.get_cmap("Blues")
    for g, ns in grp.items():
        x, yy = xy(g)
        part = min(1.0, sum(poids.get(n, 0) for n in ns) / tot) if tot else None
        coul = "#1F4E79" if part is None else ("#BFBFBF" if part == 0 else cmap(0.3 + 0.7 * part ** 0.5))
        ax.scatter([x], [yy], s=18 + 14 * math.sqrt(len(ns)), color=coul, edgecolors="white", linewidths=0.5, zorder=3)
    chaine = [c for c in topo.get("chaine", []) if c in dist]
    if chaine:
        pts = [xy((fe(c), dist[c])) for c in chaine]
        ax.plot([p[0] for p in pts], [p[1] for p in pts], color="#ED7D31", lw=1.8, alpha=0.85, zorder=2, label="Chaîne la plus longue")
    for s in sorties:
        if s in dist:
            ax.scatter(*xy((fe(s), dist[s])), s=90, color="#1F4E79", marker="s", zorder=4)
    notes, deja = [], Counter()
    for i, po in enumerate(topo.get("points_obliges", [])[:8], 1):
        c = po["cellule"]
        if c in dist:
            x, yy = xy((fe(c), dist[c]))
            ax.scatter([x], [yy], marker="*", s=120, color="#C00000", zorder=4, label="Point de passage obligé" if i == 1 else None)
            ax.annotate(str(i), (x, yy), xytext=(0, 7 + 8 * deja[(x, yy)]), textcoords="offset points", ha="center", fontsize=6.5,
                        color="#C00000", weight="bold", zorder=5)
            deja[(x, yy)] += 1
            notes.append(f"{i}. {c} {str(k.etiquette(*pos(c))[0] or '')[:34]} ({po['part']:.0%} des hypothèses)")
    if notes:
        ax.text(1.0, -0.12, "\n".join(notes), transform=ax.transAxes, ha="right", va="top", fontsize=6.3, color="#C00000")
    ax.set_xlim(-0.9, dmax + 0.9)
    ax.set_ylim(-len(bandes) + 0.5, 0.5)
    ax.set_yticks([])
    ax.set_xticks(range(0, dmax + 1))
    ax.set_xticklabels([str(dmax - i) for i in range(dmax + 1)], fontsize=6.5)
    ax.set_xlabel("Formules à traverser jusqu'à la sortie (0 = sortie)", fontsize=7.5)
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.set_title("Couloirs du calcul : une bande par feuille, de la saisie (à gauche) à la sortie (à droite)"
                 + (" ; couleur : effet mesuré, gris : sans effet" if tot else ""), fontsize=9, loc="left")
    if ax.get_legend_handles_labels()[0]:
        ax.legend(loc="upper right", fontsize=6.5, frameon=False)
    f2 = d / "carte_couloirs.png"
    fig.savefig(f2, dpi=160, bbox_inches="tight")
    plt.close(fig)
    return [str(f1), str(f2)]


def analyser(chemin, specs=(), effets=False, avec_etats=True, cache=False, cartes=None):
    t0 = time.time()
    env, k = enveloppe(chemin), Classeur(chemin)
    G, E = k.graphe()
    blocs = k.plan()
    libs = k.libelles()
    sorties = k.sorties(list(specs))
    approx = not sorties
    if approx:
        sorties = [n for n in G if G.nodes[n].get("genre") == "formule" and G.out_degree(n) == 0]
    cone_d, cone_e = set(sorties), set(sorties)
    for s in sorties:
        cone_d |= nx.ancestors(G, s) if s in G else set()
        cone_e |= nx.ancestors(E, s) if s in E else set()
    genre = lambda n: G.nodes[n].get("genre") if n in G else None
    numerique = lambda n: genre(n) == "nombre" or (genre(n) == "formule" and nombre(k.valeur(n)) is not None)
    # poids statique : cellules du cône déclaré absentes du cône effectif ; structurelles si aucune formule ne peut les lire
    decl = defaultdict(list)
    for s in k.sel:
        for b in s["boites"]:
            decl[b].append(s["structurel"])

    def structurelle(n):
        f, l, c = pos(n)
        S = [st for (fb, (c1, l1, c2, l2)), sts in decl.items() if fb == f and c1 <= c <= c2 and l1 <= l <= l2 for st in sts]
        return n in k.morts_struct or (bool(S) and all(n in st for st in S))
    sans_poids = sorted(n for n in cone_d - cone_e if numerique(n))
    s_struct = [n for n in sans_poids if structurelle(n)]
    an = lambda x: isinstance(x, int) and 1950 <= x <= 2149

    def en_tete_temps(n):
        f, l, c = pos(n)
        ws = k.f[f]
        return an(lire(ws, l, c)) and (an(lire(ws, l, c - 1)) or an(lire(ws, l, c + 1)))
    orphelins = sorted(n for n in G if genre(n) == "nombre" and G.out_degree(n) == 0 and not en_tete_temps(n))
    dans_cone = Counter(n.rsplit("!", 1)[0] for n in cone_d if genre(n) in ("nombre", "formule", "texte"))
    ilots = [f["nom"] + (f" ({f['etat']})" if f["etat"] != "visible" else "") for f in env["feuilles"]
             if f["nom"] not in dans_cone and f["nom"] not in ("Map", "Tests")]
    k.renvois(G)
    k.constantes(G, cone_d)
    sels = k.domaines(G)
    # graphe × sens : la cellule citée porte une autre période que la cellule qui cite
    for n in (x for x in cone_d if genre(x) == "formule"):
        pn, ln = k.periodes(n)
        for p in (G.predecessors(n) if pn else []):
            if genre(p) not in ("nombre", "formule") or (pos(p)[0] == pos(n)[0] and pos(p)[1] == pos(n)[1]):
                continue
            pp, lp = k.periodes(p)
            if pp and not pn & pp and not RE_LAG.search(f"{ln or ''} {lp or ''}"):
                k.ano("graphe×sens", "discordance de période", n,
                      f"cite {p} (périodes {sorted(pp)}, libellé {lp!r}) depuis une autre période", "majeure")
    # plan × graphe : sous-total (SOMME d'une sous-plage) inclus dans une SOMME plus large
    fx = lambda n: formule(lire(k.f[pos(n)[0]], pos(n)[1], pos(n)[2])) or ""
    for r in (x for x in G if genre(x) == "plage"):
        f, (c1, l1, c2, l2) = r.rsplit("!", 1)[0], range_boundaries(r.rsplit("!", 1)[1])
        ext = [s for s in G.successors(r) if genre(s) == "formule" and RE_SOMME.match(fx(s))]
        if not ext:
            continue
        for m in (x for x in G.predecessors(r) if genre(x) == "formule" and RE_SOMME.match(fx(x))):
            if all(fx(s).upper().startswith("=SUBTOTAL") for s in ext) and fx(m).upper().startswith("=SUBTOTAL"):
                continue                                         # SOUS.TOTAL ignore les SOUS.TOTAL imbriqués
            for q in (x for x in G.predecessors(m) if genre(x) == "plage" and x.rsplit("!", 1)[0] == f):
                d1, e1, d2, e2 = range_boundaries(q.rsplit("!", 1)[1])
                if c1 <= d1 and d2 <= c2 and l1 <= e1 and e2 <= l2:
                    k.ano("plan×graphe", "sous-total inclus dans une somme", m,
                          "somme d'une sous-plage, elle-même incluse dans une somme plus large : double compte", "critique")
    # manifeste pré-rempli : hypothèses lues directement (les tables lues par plage sont des données, pas des inputs)
    carte, vus, lignes, chk = [], Counter(), defaultdict(list), defaultdict(list)

    def cle(prefixe, lib):
        c = f"{prefixe}.{slug(lib)}"
        vus[c] += 1
        return c if vus[c] == 1 else f"{c}_{vus[c]}"

    def segments(dd):
        for (f, x), vals in sorted(dd.items()):
            seg = []
            for v in sorted(vals):
                if seg and v == seg[-1][1] + 1:
                    seg[-1][1] = v
                else:
                    seg.append([v, v])
            for a, b in seg:
                yield f, x, a, b
    for n in cone_d:
        if genre(n) == "nombre" and n not in sels and n not in k.durs and any(genre(s) == "formule" for s in G.successors(n)):
            f, l, c = pos(n)
            lignes[(f, l)].append(c)
    for n in G:                                       # un check calcule un écart ou une comparaison, pas un simple lien
        f = pos(n)[0]
        if genre(n) == "formule" and not RE_DOC.search(f) and RE_CHK.search(str(k.periodes(n)[1] or "")) and \
                re.search(r"[^=<>]-|<>|[<>]=?|(?<=.)=|\bABS\(", fx(n)[1:]):
            chk[(f, pos(n)[1])].append(pos(n)[2])
    fmt = lambda v: str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)
    for n, s in sorted(sels.items()):                 # sélecteurs globaux ou à liste ; les règles de ligne restent dans le JSON
        if s["formules"] >= 3 or "liste" in s["origines"]:
            f, l, c = pos(n)
            carte.append({"role": "selecteur_scenario", "cle": cle("sel", k.etiquette(f, l, c)[0] or k.etiquette(f, l, c)[1]),
                          "feuille": f, "plage": f"{GL(c)}{l}", "attendu": ";".join(dict.fromkeys(map(fmt, s["valeurs"])))})
    for f, l, a, b in segments(lignes):
        lib = k.etiquette(f, l, a)[0]
        carte.append({"role": "input", "cle": cle("inp", lib), "feuille": f, "plage": f"{GL(a)}{l}" + (f":{GL(b)}{l}" if b > a else ""),
                      "libelle": lib, "perturber": "non" if re.search(r"\bdate|ann[ée]e", str(lib), re.I) else ""})
    carte += [{"role": "output", "cle": cle("out", s), "feuille": s.rpartition("!")[0].strip("'") or None,
               "plage": s.rpartition("!")[2]} for s in specs]
    for f, l, a, b in segments(chk):
        lib = k.etiquette(f, l, a)[0]
        carte.append({"role": "check", "cle": cle("chk", lib), "feuille": f, "plage": f"{GL(a)}{l}" + (f":{GL(b)}{l}" if b > a else ""),
                      "libelle": lib, "a_valider": True})
    textes = sorted(n for n in cone_d if genre(n) == "texte" or (genre(n) == "formule" and isinstance(k.valeur(n), str)))
    res = {"enveloppe": env,
           "plan": {"formules": len(k.sig), "blocs": len(blocs), "blocs_principaux": blocs[:15]},
           "graphe": {"noeuds": G.number_of_nodes(), "aretes": G.number_of_edges(), "opaques": dict(k.opaques),
                      "sorties": sorties, "sorties_approximatives": approx, "cone_declare": len(cone_d),
                      "cone_effectif": len(cone_e), "cellules_numeriques_du_cone": sum(map(numerique, cone_d - set(sorties))),
                      "textes_du_cone": len(textes), "feuilles_du_cone": dict(dans_cone), "ilots": ilots,
                      "constantes_orphelines": len(orphelins), "exemples_orphelines": orphelins[:15]},
           "sens": {"libelles": len(libs), "par_origine": dict(Counter(x["origine"] for x in libs)), "formules_cache": env["formules_cache"]},
           "poids": {"appels_selectifs_evalues": len(k.sel), "sans_poids": len(sans_poids), "dont_structurels": len(s_struct),
                     "par_colonne": dict(Counter(n.rsplit("!", 1)[0] + "!" + re.sub(r"\d", "", n.rsplit("!", 1)[1])
                                                 for n in sans_poids).most_common(12)),
                     "exemples_sans_poids": sans_poids[:20], "exemples_structurels": s_struct[:20]},
           "selecteurs": {n: {**s, "valeurs": s["valeurs"][:20]} for n, s in sels.items()},
           "constantes_a_qualifier": [{"cellule": n, "motif": m, "libelle": k.etiquette(*pos(n))[0], "valeur": k.valeur(n),
                                       "dans_cone": n in cone_d} for n, m in sorted(k.durs.items(), key=lambda x: (x[0] not in cone_d, x[0]))],
           "manifeste_propose": carte}
    poids = eff_tout = m = None
    if effets or cache:
        try:
            m = Moteur(chemin)
        except Exception as ex:                                  # LibreOffice ou python3-uno absents
            res["moteur_indisponible"] = str(ex)[:300]
    if m:
        try:
            formules = sorted(n for n in cone_d if genre(n) == "formule" and n not in k.matricielles)
            ec = [{"cellule": n, "cache": k.valeur(n), "recalcul": m.valeur(n)} for n in formules if not proche(k.valeur(n), m.valeur(n))]
            res["cache"] = {"moteur": "LibreOffice", "formules_comparees": len(formules), "ecarts": len(ec), "exemples": ec[:20],
                            "sorties_touchees": [e["cellule"] for e in ec if e["cellule"] in sorties]}
            for e in ec[:200]:
                k.ano("cache×recalcul", "cache périmé ou calcul divergent", e["cellule"], "valeur en cache différente du recalcul LibreOffice",
                      "critique" if e["cellule"] in sorties else "majeure")
            if effets:
                cibles = sorted(n for n in cone_d - set(sorties) if genre(n) in ("nombre", "formule", "texte") and n not in k.matricielles)
                t1 = time.time()
                mes, base, derive = mesurer(m, cibles, sorties)
                num = {n: r for n, r in mes.items() if not isinstance(r["valeur"], str)}
                nul = sorted(n for n, r in num.items() if r["classe"] == "sans effet")
                base_vals = {n: m.valeur(n) for n in formules}
                absorbe = {n: absorbeurs(m, G, cone_d, n, base_vals) for n in nul[:800]}
                parabs = Counter(a for v in absorbe.values() for a in v)
                eff = lambda r: max((abs(x) for x in r.get("delta", []) if not isinstance(x, str)), default=0.0)
                eff_tout = {n: eff(r) for n, r in mes.items() if not isinstance(r["valeur"], str)}   # +1 % : nombres seulement
                poids = {n: v for n, v in eff_tout.items() if genre(n) == "nombre"}
                tot = sum(poids.values()) or 1
                parf = Counter()
                for n, p in poids.items():
                    parf[n.rsplit("!", 1)[0]] += p
                lib = lambda n: " / ".join(str(x) for x in k.etiquette(*pos(n)) if x)
                res["effets"] = {
                    "base_libreoffice": base, "derive_finale": derive, "duree_s": round(time.time() - t1, 1),
                    "testees": len(mes), "numeriques": len(num), "textes": len(mes) - len(num),
                    "classes": dict(Counter(r["classe"] for r in mes.values())),
                    "classes_numeriques": dict(Counter(r["classe"] for r in num.values())),
                    "sans_effet": len(nul), "dont_vues_par_les_masques": len(set(nul) & set(sans_poids)),
                    "masques_dementis": sorted(set(sans_poids) - set(nul) - set(k.matricielles))[:30],
                    "sans_effet_par_colonne": dict(Counter(n.rsplit("!", 1)[0] + "!" + re.sub(r"\d", "", n.rsplit("!", 1)[1])
                                                           for n in nul).most_common(15)),
                    "principaux": [{"cellule": n, "libelle": lib(n), "valeur": mes[n]["valeur"], "delta_1pct": mes[n]["delta"]}
                                   for n in sorted(poids, key=lambda x: -poids[x])[:25] if poids[n] > 0],
                    "par_feuille": {f: round(p / tot, 3) for f, p in parf.most_common()},
                    "seuils": [{"cellule": n, "libelle": lib(n), **r["seuil"]} for n, r in mes.items() if r["classe"] == "seuil"][:40],
                    "textes_pesants": [{"cellule": n, "valeur": r["valeur"], "delta": r["delta"]} for n, r in mes.items()
                                       if isinstance(r["valeur"], str) and r["classe"] != "sans effet"][:40],
                    "absorbeurs": [{"formule": a, "texte": fx(a), "absorbe": c, "exemples": [n for n, v in absorbe.items() if a in v][:5]}
                                   for a, c in parabs.most_common(15)],
                    "detail": {n: {**r, "absorbeurs": absorbe.get(n, [])[:5]} for n, r in mes.items()}}
                if avec_etats and sels:
                    et = etats(m, sels, sorties, base, nul)
                    res["etats"] = et
                    rev = {x for e in et for x in e["reveillees"]}
                    res["effets"]["sans_effet_dans_tous_les_etats"] = len(set(nul) - rev)
        finally:
            m.fermer()
    topo = topologie(G, sorties, cone_d, poids)
    res["topologie"] = {x: v for x, v in topo.items() if not x.startswith("_")}
    for p in res["topologie"].get("points_obliges", []):
        p["libelle"] = k.etiquette(*pos(p["cellule"]))[0]
    if cartes:
        res["cartes"] = dessiner(k, G, cone_d, sorties, env, topo, eff_tout, cartes)
    reg = regrouper(k.anos, cone_d)
    res["anomalies"] = [a for a in reg if a["dans_cone"] or a["severite"] != "à vérifier"]
    res["alertes_mineures_hors_cone"] = len(reg) - len(res["anomalies"])
    res["duree_s"] = round(time.time() - t0, 1)
    return res, libs


def rapport_md(res):
    e, p, g, s = res["enveloppe"], res["poids"], res["graphe"], res["sens"]
    fc = s["formules_cache"]
    L = ["# Analyse fine", "",
         f"- Enveloppe : {len(e['feuilles'])} feuilles dont {sum(f['etat'] != 'visible' for f in e['feuilles'])} masquée(s) ; "
         f"types {e['types_cellules']} ; noms {e['noms_definis']} ; liens externes {e['liens_externes']} ; "
         f"DUMMYFUNCTION {e['dummyfunction']} ; macros {e['macros']} ; Power Query {e['power_query']}",
         f"- Cache des formules : {fc.get('valeur', 0)} avec valeur, {fc.get('chaine_vide', 0)} renvoyant une chaîne vide, "
         f"{fc.get('sans_valeur', 0)} sans valeur" + (" : recalculer une copie avant toute lecture" if fc.get("sans_valeur") else ""),
         f"- Plan : {res['plan']['formules']} formules en {res['plan']['blocs']} blocs",
         f"- Graphe : {g['noeuds']} nœuds, {g['aretes']} arêtes, références opaques {g['opaques'] or 'aucune'} ; cône déclaré {g['cone_declare']}, "
         f"effectif {g['cone_effectif']} ; {g['cellules_numeriques_du_cone']} cellules numériques et {g['textes_du_cone']} textes en amont"
         + (" (sorties non déclarées : cône approximatif)" if g["sorties_approximatives"] else "")
         + f" ; îlots : {', '.join(g['ilots']) or 'aucun'} ; constantes orphelines : {g['constantes_orphelines']}",
         f"- Sens : {s['libelles']} libellés positionnés {s['par_origine']}",
         f"- Poids lu dans les masques : {p['appels_selectifs_evalues']} appels sélectifs évalués ; {p['sans_poids']} cellules du cône "
         f"sans poids dans l'état courant, dont {p['dont_structurels']} structurelles ; par colonne {p['par_colonne']}"]
    if "moteur_indisponible" in res:
        L.append(f"- Moteur de recalcul indisponible : {res['moteur_indisponible']}")
    if "cache" in res:
        q = res["cache"]
        L.append(f"- Cache contre recalcul LibreOffice : {q['ecarts']} écart(s) sur {q['formules_comparees']} formules du cône"
                 + (f" ; sorties touchées : {q['sorties_touchees']}" if q["sorties_touchees"] else ""))
    if "effets" in res:
        q = res["effets"]
        L += ["", "## Effets mesurés (LibreOffice, une cellule à la fois)", "",
              f"- {q['testees']} cellules essayées ({q['numeriques']} numériques, {q['textes']} textes) en {q['duree_s']} s ; "
              f"classes {q['classes']} ; dérive finale {q['derive_finale']}",
              f"- Cellules numériques sans effet : {q['sans_effet']}, dont {q['dont_vues_par_les_masques']} déjà vues par les masques"
              + (f" ; sans effet dans tous les états essayés : {q['sans_effet_dans_tous_les_etats']}" if "sans_effet_dans_tous_les_etats" in q else "")
              + (f" ; masques démentis : {len(q['masques_dementis'])}" if q["masques_dementis"] else ""),
              f"- Par colonne : {q['sans_effet_par_colonne']}",
              f"- Part de l'effet par feuille : {q['par_feuille']}", "",
              "| Cellule | Libellé | Valeur | Écart de sortie pour +1 % |", "|---|---|---|---|"]
        L += [f"| {x['cellule']} | {x['libelle']} | {x['valeur']:.4g} | {', '.join(f'{d:+.4g}' if not isinstance(d, str) else d for d in x['delta_1pct'])} |"
              for x in q["principaux"][:15]]
        if q["seuils"]:
            L += ["", "Effets de seuil (rien à +1 %, un effet à forte amplitude) : "
                  + "; ".join(f"{x['cellule']} ({x['essai']})" for x in q["seuils"][:12])]
        if q["absorbeurs"]:
            L += ["", "| Absorbeur (la perturbation s'y arrête) | Formule | Cellules absorbées |", "|---|---|---|"]
            L += [f"| {a['formule']} | `{a['texte'][:70]}` | {a['absorbe']} |" for a in q["absorbeurs"][:10]]
    if res.get("etats"):
        L += ["", "## Sélecteurs essayés valeur par valeur", "", "| Sélecteur | Valeur | Écart de sortie | Cellules réveillées |", "|---|---|---|---|"]
        L += [f"| {x['selecteur']} | {x['valeur']} | {x['ecart']} | {len(x['reveillees'])} |" for x in res["etats"][:30]]
    t = res["topologie"]
    L += ["", "## Topologie", "", f"- {t['hypotheses']} hypothèses en amont ; cycles : {'oui' if t['cycles'] else 'non'}"
          + (f" ; profondeur maximale {t['profondeur_max']} formules ; chaîne la plus longue, de la saisie à la sortie : "
             f"{', '.join(t['chaine'][:8])}{', …' if len(t['chaine']) > 8 else ''}"
             if "chaine" in t else "")]
    if t.get("points_obliges"):
        L += ["", "| Point de passage obligé | Libellé | Hypothèses dominées | Part | Part de l'effet |", "|---|---|---|---|---|"]
        L += [f"| {x['cellule']} | {x.get('libelle') or ''} | {x['hypotheses']} | {x['part']:.0%} | "
              f"{'' if x['part_effet'] is None else format(x['part_effet'], '.0%')} |" for x in t["points_obliges"][:10]]
    if res.get("cartes"):
        L += ["", "Cartes : " + ", ".join(res["cartes"])]
    cq = [x for x in res["constantes_a_qualifier"] if x["dans_cone"]]
    if res["constantes_a_qualifier"]:
        L += ["", f"## Constantes à qualifier, retirées du manifeste ({len(cq)} dans le cône, "
              f"{len(res['constantes_a_qualifier']) - len(cq)} hors cône, détail dans le JSON)", "",
              "| Cellule | Motif | Libellé | Valeur |", "|---|---|---|---|"]
        L += [f"| {x['cellule']} | {x['motif']} | {x['libelle']} | {x['valeur']} |" for x in cq[:40]]
    L += ["", f"## Registre d'anomalies ({res['alertes_mineures_hors_cone']} alertes « à vérifier » hors cône, dans le JSON seulement)", "",
          "| Sévérité | Cône | Croisement | Type | Cellules | Détail |", "|---|---|---|---|---|---|"]
    L += [f"| {a['severite']} | {'oui' if a['dans_cone'] else 'non'} | {a['croisement']} | {a['type']} | {a['cellule']} | {a['detail']} |"
          for a in res["anomalies"][:80]]
    return "\n".join(L)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("classeur")
    ap.add_argument("--sorties", nargs="*", default=[])
    ap.add_argument("--md")
    ap.add_argument("--json")
    ap.add_argument("--libelles")
    ap.add_argument("--effets", action="store_true")
    ap.add_argument("--sans-etats", action="store_true")
    ap.add_argument("--cartes")
    ap.add_argument("--verifier-cache", action="store_true")
    ap.add_argument("--recalculer-copie", help="recalcule une copie dans LibreOffice, l'enregistre ici et analyse cette copie")
    a = ap.parse_args()
    if a.recalculer_copie:
        m = Moteur(a.classeur)
        try:
            m.enregistrer(a.recalculer_copie)
        finally:
            m.fermer()
        a.classeur = a.recalculer_copie
    res, libs = analyser(a.classeur, a.sorties, a.effets, not a.sans_etats, a.verifier_cache, a.cartes)
    for chemin, contenu in ((a.json, res), (a.libelles, libs)):
        if chemin:
            Path(chemin).write_text(json.dumps(contenu, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    md = rapport_md(res)
    if a.md:
        Path(a.md).write_text(md, encoding="utf-8")
    print(md)


if __name__ == "__main__":
    main()
